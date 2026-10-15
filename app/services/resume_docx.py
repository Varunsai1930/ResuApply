"""Editable Word export of the same selected facts used by the resume preview."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from urllib.parse import urlsplit

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .resume_layouts import get_layout


class ResumeDocxError(ValueError):
    """The selected resume could not be represented as a Word document."""


_STYLES = {
    "classic": {"font": "Arial", "size": 10.5, "name": 22, "heading": 12,
                "leading": 1.40, "after": 3, "section": 9,
                "accent": "222222"},
    "modern": {"font": "Calibri", "size": 10.5, "name": 24, "heading": 11,
               "leading": 1.45, "after": 3, "section": 10,
               "accent": "1F4E79"},
    "compact": {"font": "Arial", "size": 10, "name": 19, "heading": 10.5,
                "leading": 1.32, "after": 2, "section": 7,
                "accent": "222222"},
}


def _text(value: object, field: str) -> str:
    if not isinstance(value, str):
        raise ResumeDocxError(f"Resume {field} must be text. Review the selected resume before exporting.")
    # XML 1.0 cannot represent these characters. Never silently drop accepted text.
    if any(not (char in "\t\n\r" or 0x20 <= ord(char) <= 0xD7FF or
                0xE000 <= ord(char) <= 0xFFFD or 0x10000 <= ord(char) <= 0x10FFFF)
           for char in value):
        raise ResumeDocxError(f"Resume {field} contains characters Word cannot read. Correct the text before exporting.")
    return value


def _mapping(value: object, field: str) -> dict:
    if not isinstance(value, dict):
        raise ResumeDocxError(f"Resume {field} is damaged. Review the selected resume before exporting.")
    return value


def _items(value: object, field: str) -> list:
    if not isinstance(value, list):
        raise ResumeDocxError(f"Resume {field} is damaged. Review the selected resume before exporting.")
    return value


def _field(value: dict, key: str) -> str:
    return _text(value.get(key, ""), key)


def _dates(entry: dict) -> str:
    start, end = _field(entry, "start"), _field(entry, "end")
    if end == "present":
        end = "Present"
    return " – ".join(value for value in (start, end) if value)


def _web_url(value: str) -> bool:
    if any(char.isspace() or ord(char) < 0x20 for char in value):
        return False
    try:
        parsed = urlsplit(value)
        parsed.port  # Invalid/out-of-range port numbers are not usable web URLs.
        return parsed.scheme.lower() in ("http", "https") and bool(parsed.hostname)
    except ValueError:
        return False


def _style(document, name: str, *, font: str, size: float, bold: bool = False,
           color: str | None = None):
    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles["Normal"]
    style.font.name = font
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    return style


def _configure(document, key: str) -> None:
    values = _STYLES[key]
    for section in document.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(15)
        section.left_margin = section.right_margin = Mm(17)
    normal = document.styles["Normal"]
    normal.font.name = values["font"]
    normal.font.size = Pt(values["size"])
    normal.paragraph_format.line_spacing = values["leading"]
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(values["after"])
    normal.paragraph_format.keep_together = False
    normal.paragraph_format.keep_with_next = False
    normal.paragraph_format.widow_control = True

    title = document.styles["Title"]
    title.base_style = normal
    title.font.name = values["font"]
    title.font.size = Pt(values["name"])
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(5 if key == "compact" else 7)
    title.paragraph_format.line_spacing = 1.15
    title.paragraph_format.keep_with_next = False
    _remove_borders(title.element.get_or_add_pPr())
    heading = _style(document, "Resume Section", font=values["font"], size=values["heading"],
                     bold=True, color=values["accent"])
    heading.paragraph_format.space_before = Pt(values["section"])
    heading.paragraph_format.space_after = Pt(4 if key != "compact" else 3)
    heading.paragraph_format.keep_with_next = True
    _rule(heading.element.get_or_add_pPr(), "B9C7D5" if key == "modern" else "555555")
    _style(document, "Resume Entry", font=values["font"], size=values["size"], bold=True)
    _style(document, "Resume Metadata", font=values["font"], size=values["size"])
    bullet = _style(document, "Resume Bullet", font=values["font"], size=values["size"])
    bullet.paragraph_format.space_after = Pt(values["after"])
    hyperlink = document.styles.add_style("Resume Hyperlink", WD_STYLE_TYPE.CHARACTER)
    hyperlink.font.name = values["font"]
    hyperlink.font.size = Pt(values["size"])
    hyperlink.font.color.rgb = RGBColor.from_string(values["accent"])
    hyperlink.font.underline = True

    # Do not inherit template metadata or include job details/source IDs in the file.
    props = document.core_properties
    for name in ("title", "subject", "author", "keywords", "comments", "last_modified_by",
                 "category", "content_status", "identifier", "language", "version"):
        setattr(props, name, "")
    props.revision = 1
    props.created = props.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)


def _rule(properties, color: str, *, size: str = "4") -> None:
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", size), ("space", "4"), ("color", color)):
        bottom.set(qn("w:" + key), value)
    borders.append(bottom)
    properties.append(borders)


def _remove_borders(properties) -> None:
    # Word's built-in Title can carry a template-specific blue underline.
    for border in properties.findall(qn("w:pBdr")):
        properties.remove(border)


def _bullet_numbering(document) -> int:
    numbering = document.part.numbering_part.element
    abstracts = [int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))]
    numbers = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    abstract_id, num_id = max(abstracts, default=-1) + 1, max(numbers, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level_type = OxmlElement("w:multiLevelType")
    level_type.set(qn("w:val"), "singleLevel")
    abstract.append(level_type)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("start", "1"), ("numFmt", "bullet"), ("lvlText", "•"), ("lvlJc", "left")):
        child = OxmlElement(f"w:{tag}")
        child.set(qn("w:val"), value)
        level.append(child)
    paragraph = OxmlElement("w:pPr")
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "240")
    indent.set(qn("w:hanging"), "180")
    paragraph.append(indent)
    level.append(paragraph)
    abstract.append(level)
    # Abstract definitions precede concrete numbering instances in OOXML.
    numbering.insert(0, abstract)
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), str(abstract_id))
    number.append(reference)
    numbering.append(number)
    return num_id


def _bullet(document, value: str, num_id: int) -> None:
    paragraph = document.add_paragraph(value, "Resume Bullet")
    number = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    number.get_or_add_ilvl().val = 0
    number.get_or_add_numId().val = num_id


def _link(document, value: str, *, keep_next: bool = False) -> None:
    paragraph = document.add_paragraph(style="Resume Metadata")
    paragraph.paragraph_format.keep_with_next = keep_next
    if not _web_url(value):
        paragraph.add_run(value)
        return
    from docx.opc.constants import RELATIONSHIP_TYPE

    relationship = document.part.relate_to(value, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "ResumeHyperlink")
    properties.append(style)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = value
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _entry(document, heading: str, metadata: list[tuple[str, bool]], bullets: list[str], num_id: int) -> None:
    paragraph = document.add_paragraph(heading, "Resume Entry")
    paragraph.paragraph_format.keep_with_next = bool(metadata or bullets)
    for index, (text, link) in enumerate(metadata):
        keep_next = index < len(metadata) - 1 or bool(bullets)
        if link:
            _link(document, text, keep_next=keep_next)
        else:
            paragraph = document.add_paragraph(text, "Resume Metadata")
            paragraph.paragraph_format.keep_with_next = keep_next
    for text in bullets:
        _bullet(document, text, num_id)


def export_resume(rendered: dict, layout: str = "classic") -> bytes:
    """Render trusted ``resume.render_context`` facts without expanding selection."""
    try:
        key = get_layout(layout).key
    except (TypeError, ValueError) as exc:
        raise ResumeDocxError("Choose a valid resume layout before exporting.") from exc
    rendered = _mapping(rendered, "content")
    contact = _mapping(rendered.get("contact"), "contact")
    document = Document()
    _configure(document, key)
    num_id = _bullet_numbering(document)
    title_paragraph = document.add_paragraph(_field(contact, "name"), "Title")
    _remove_borders(title_paragraph._p.get_or_add_pPr())
    details = [_field(contact, key) for key in ("email", "phone", "location")]
    if any(details):
        document.add_paragraph(" · ".join(value for value in details if value))
    links = _mapping(contact.get("links", {}), "contact links")
    for key in ("linkedin", "github", "portfolio"):
        value = _field(links, key)
        if value:
            _link(document, value)
    if layout == "modern" and len(document.paragraphs) > 1:
        _rule(document.paragraphs[-1]._p.get_or_add_pPr(), "1F4E79", size="12")

    summary = rendered.get("summary")
    if summary is not None:
        document.add_paragraph("Summary", "Resume Section")
        document.add_paragraph(_field(_mapping(summary, "summary"), "text"))

    for section, title in (("experience", "Experience"), ("projects", "Projects")):
        entries = _items(rendered.get(section, []), section)
        if entries:
            document.add_paragraph(title, "Resume Section")
        for raw in entries:
            entry = _mapping(raw, section)
            if section == "experience":
                heading = " — ".join(value for value in (_field(entry, "title"), _field(entry, "organization")) if value)
            else:
                heading = " — ".join(value for value in (_field(entry, "name"), _field(entry, "role")) if value)
            metadata = [(value, False) for value in (_field(entry, "location"), _dates(entry)) if value]
            link = _field(entry, "link")
            if link:
                metadata.append((link, True))
            technologies = [_text(value, "technology") for value in _items(entry.get("technologies", []), "technologies")]
            if technologies:
                metadata.append(("Technologies: " + ", ".join(technologies), False))
            bullets = [_field(_mapping(value, "bullet"), "text")
                       for value in _items(entry.get("bullets", []), "bullets")]
            _entry(document, heading, metadata, bullets, num_id)

    education = _items(rendered.get("education", []), "education")
    if education:
        document.add_paragraph("Education", "Resume Section")
    for raw in education:
        entry = _mapping(raw, "education")
        degree = ", ".join(value for value in (_field(entry, "degree"), _field(entry, "field")) if value)
        metadata = [(value, False) for value in (degree, _dates(entry)) if value]
        gpa = _field(entry, "gpa")
        if gpa:
            metadata.append(("GPA: " + gpa, False))
        bullets = [_field(_mapping(value, "education bullet"), "text")
                   for value in _items(entry.get("bullets", []), "education bullets")]
        _entry(document, _field(entry, "institution"), metadata, bullets, num_id)

    skills = [_text(value, "skill") for value in _items(rendered.get("skills", []), "skills")]
    if skills:
        document.add_paragraph("Skills", "Resume Section")
        document.add_paragraph(", ".join(skills))
    certifications = _items(rendered.get("certifications", []), "certifications")
    if certifications:
        document.add_paragraph("Certifications", "Resume Section")
    for raw in certifications:
        entry = _mapping(raw, "certification")
        value = " — ".join(text for text in (_field(entry, "name"), _field(entry, "issuer")) if text)
        date = _field(entry, "date")
        if date:
            value += f" ({date})"
        _bullet(document, value, num_id)
    result = BytesIO()
    document.save(result)
    return result.getvalue()
