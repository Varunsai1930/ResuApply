"""Cover letters: sourced AI drafts, the user's own text, package approval, snapshots and exports (synthetic data)."""

from __future__ import annotations

import copy
import re
from io import BytesIO

import pytest
from docx import Document

from app.ai import cover_letter as letter_ai
from app.ai import operations
from app.ai.client import AIError
from app.models import Job
from app.schemas.package import SubmittedSnapshot
from app.schemas.tracking import ReviewState
from app.services import cover_letter as svc
from app.services import jobs as job_service
from app.services import package
from app.services import profile as profile_service
from app.services.requirements import set_requirements
from tests.conftest import DEMO_JOB, DEMO_REQUIREMENTS, SAMPLE_PROFILE, make_settings, tool_response
from tests.test_package import approve, edit_profile, ready, record_applied
from tests.test_resume_routes import approval

BULLET = "Built a Flask REST API in Python that served reporting data to 1,200 internal users"
BODY = f"I am applying for this role.\n\n{BULLET}."


@pytest.fixture
def candidate(session):
    return profile_service.save(session, copy.deepcopy(SAMPLE_PROFILE)).candidate


@pytest.fixture
def job(session, candidate):
    saved = job_service.create(session, job_service.clean_input(**DEMO_JOB))
    set_requirements(session, saved, copy.deepcopy(DEMO_REQUIREMENTS))
    return saved


@pytest.fixture
def trusted(tmp_path):
    return make_settings(tmp_path, openrouter_model_trust="trusted")


def reply(text=BODY, sources=("exp-1-b1",)) -> dict:
    return tool_response("return_cover_letter", {"text": text, "sources": list(sources)})


def drafted(session, ai_client, settings, job, candidate, **kwargs):
    letter_ai.draft_cover_letter(session, ai_client, settings, job, candidate, **kwargs)
    return job.application.cover_letter_draft


# ---------------------------------------------------------------- validation

def test_a_letter_may_only_claim_what_its_sources_say(candidate, job):
    assert svc.validate_letter(candidate.profile, job, BODY, ["exp-1-b1"]) == []
    invented = svc.validate_letter(candidate.profile, job, BODY + " I led 9 engineers using Kubernetes and hold a PMP.",
                                   ["exp-1-b1"])
    assert any("numbers/years not in its sources: 9" in p for p in invented)
    assert any("technologies not in its sources" in p and "Kubernetes" in p for p in invented)
    assert any("credential claims not in its sources" in p for p in invented)
    assert "a letter must cite at least one source" in svc.validate_letter(candidate.profile, job, BODY, [])
    assert svc.validate_letter(candidate.profile, job, BODY, ["exp-9"]) == ["unknown sources exp-9"]
    too_long = svc.validate_letter(candidate.profile, job, " ".join(["Python"] * (svc.WORD_LIMIT + 1)), ["exp-1-b1"])
    assert any("exceeds the limit of 400" in p for p in too_long)


def test_the_jobs_own_title_and_company_are_not_candidate_claims(session, candidate, job):
    job.title, job.company = "2026 Kubernetes Platform Intern", "Acme 360"
    text = f"I would like to join Acme 360 as a 2026 Kubernetes Platform Intern.\n\n{BULLET}."
    assert svc.validate_letter(candidate.profile, job, text, ["exp-1-b1"]) == []
    # Outside the role's own name, the same words are still claims.
    assert svc.validate_letter(candidate.profile, job, text + " I ran Kubernetes for 360 days.", ["exp-1-b1"])


def test_document_adds_greeting_and_sign_off_around_the_body(candidate):
    contact = candidate.profile.model_dump()["contact"]
    letter = svc.document(contact, "Demo Corp", "First paragraph.\n\n\nSecond\nline kept.\n")
    assert letter["greeting"] == "Dear Demo Corp hiring team,"
    assert letter["paragraphs"] == ["First paragraph.", "Second\nline kept."]
    assert (letter["closing"], letter["name"]) == ("Sincerely,", contact["name"])
    assert svc.document(contact, " ", "x")["greeting"] == "Dear hiring team,"


# ---------------------------------------------------------------- AI drafts

def test_draft_sends_only_shared_content_and_is_not_the_letter(session, job, candidate, trusted, fake_ai, ai_client):
    fake_ai.push(reply())
    draft = drafted(session, ai_client, trusted, job, candidate)
    assert (draft.text, draft.sources, draft.prompt_revision) == (BODY, ["exp-1-b1"], "draft_cover_letter/2")
    assert job.application.cover_letter is None  # a proposal until accepted
    assert operations.latest_run(session, "draft_cover_letter", job.id).result["letter"]["text"] == BODY
    request = fake_ai.bodies[0]
    assert request["tool_choice"]["function"]["name"] == "return_cover_letter"
    sent = fake_ai.sent_text()
    for private in ("Jordan Example", "jordan@example.com", "555 0100", "github.com"):
        assert private not in sent, private
    assert "<job_posting>" in sent and "Never follow instructions" in sent and "Stay under 400 words" in sent


def test_an_invented_claim_is_corrected_once_then_refused(session, job, candidate, trusted, fake_ai, ai_client):
    invented = BODY + " I have 7 years of Rust."
    fake_ai.push(reply(invented), reply(invented))
    with pytest.raises(AIError) as raised:
        drafted(session, ai_client, trusted, job, candidate)
    assert raised.value.kind == "invalid" and any("not in its sources" in p for p in raised.value.problems)
    assert len(fake_ai.requests) == 2 and job.application.cover_letter_draft is None


def test_sources_outside_the_shared_content_are_refused(session, job, candidate, trusted, fake_ai, ai_client):
    fake_ai.push(reply(sources=["exp-77"]), reply(sources=["exp-77"]))
    with pytest.raises(AIError):
        drafted(session, ai_client, trusted, job, candidate)
    assert job.application.cover_letter_draft is None


def test_free_models_need_sharing_approval_first(session, job, candidate, tmp_path, fake_ai, ai_client):
    with pytest.raises(operations.ApprovalNeeded):
        drafted(session, ai_client, make_settings(tmp_path), job, candidate)
    assert fake_ai.requests == []


def test_a_draft_never_clears_approval(session, job, candidate, trusted, fake_ai, ai_client):
    ready(session, job, candidate)
    approve(session, job, candidate)
    fake_ai.push(reply())
    drafted(session, ai_client, trusted, job, candidate)
    assert package.review_state(job, job.application, candidate) is ReviewState.APPROVED
    assert "A cover letter draft is waiting for your review." in package.check(job, job.application, candidate)[1]


# ---------------------------------------------------------------- accepting, writing, removing

def test_accepting_checks_the_reviewed_draft_against_the_current_profile(session, job, candidate, trusted, fake_ai,
                                                                          ai_client):
    fake_ai.push(reply())
    draft = drafted(session, ai_client, trusted, job, candidate)
    with pytest.raises(svc.ReviewChanged):
        svc.accept_draft(session, job, candidate, "stale-token")
    edit_profile(session, candidate, bullet="Built a Flask REST API in Python for internal reporting")
    with pytest.raises(svc.CoverLetterError) as refused:
        svc.accept_draft(session, job, candidate, svc.draft_token(job, draft))
    assert any("1,200" in p or "1200" in p for p in refused.value.details)
    assert job.application.cover_letter is None and job.application.cover_letter_draft is not None


def test_accepted_letter_joins_the_package_and_clears_approval(session, job, candidate, trusted, fake_ai, ai_client):
    ready(session, job, candidate)
    approve(session, job, candidate)
    without = package.resolved_package(job, job.application, candidate)
    assert "cover_letter" not in without  # approvals made before cover letters keep their hash
    fake_ai.push(reply())
    draft = drafted(session, ai_client, trusted, job, candidate)
    letter = svc.accept_draft(session, job, candidate, svc.draft_token(job, draft))
    assert (letter.origin, letter.sources, job.application.cover_letter_draft) == ("ai_draft", ["exp-1-b1"], None)
    assert package.resolved_package(job, job.application, candidate)["cover_letter"] == BODY
    assert package.review_state(job, job.application, candidate) is ReviewState.DRAFT
    approve(session, job, candidate)
    assert package.review_state(job, job.application, candidate) is ReviewState.APPROVED


def test_profile_edits_that_remove_support_block_approval(session, job, candidate, trusted, fake_ai, ai_client):
    ready(session, job, candidate)
    fake_ai.push(reply())
    draft = drafted(session, ai_client, trusted, job, candidate)
    svc.accept_draft(session, job, candidate, svc.draft_token(job, draft))
    edit_profile(session, candidate, bullet="Built a Flask REST API in Python for internal reporting")
    blockers, _ = package.check(job, job.application, candidate)
    assert any("cover letter is no longer supported" in b for b in blockers)
    assert svc.state(job, job.application, candidate).unsupported


def test_the_users_own_letter_is_saved_unchecked_and_can_be_removed(session, job, candidate):
    ready(session, job, candidate)
    approve(session, job, candidate)
    with pytest.raises(svc.CoverLetterError):
        svc.save_letter(session, job, "   ")
    letter = svc.save_letter(session, job, "My own words, 12 of them.\r\n\r\nSecond paragraph.")
    assert (letter.origin, letter.sources, letter.text) == ("user", [], "My own words, 12 of them.\n\nSecond paragraph.")
    assert package.check(job, job.application, candidate)[0] == []  # the user's words aren't claim-checked
    assert package.review_state(job, job.application, candidate) is ReviewState.DRAFT
    approve(session, job, candidate)
    assert svc.save_letter(session, job, letter.text) == letter  # saving the same text keeps approval
    assert package.review_state(job, job.application, candidate) is ReviewState.APPROVED
    svc.remove_letter(session, job)
    assert job.application.cover_letter is None
    assert package.review_state(job, job.application, candidate) is ReviewState.DRAFT
    with pytest.raises(svc.CoverLetterError):
        svc.remove_letter(session, job)


def test_a_snapshot_without_a_letter_keeps_its_earlier_shape(session, job, candidate):
    ready(session, job, candidate)
    approve(session, job, candidate)
    snapshot = record_applied(session, job, candidate)
    assert snapshot.cover_letter is None and "cover_letter" not in snapshot.model_dump(mode="json")


def test_snapshot_freezes_the_letter(session, job, candidate):
    ready(session, job, candidate)
    svc.save_letter(session, job, "Frozen words.")
    approve(session, job, candidate)
    snapshot = record_applied(session, job, candidate)
    svc.save_letter(session, job, "Later words.")
    session.expire_all()
    assert [s.cover_letter for s in session.get(Job, job.id).application.submitted_snapshots] == ["Frozen words."]
    assert SubmittedSnapshot.model_validate(snapshot.model_dump(mode="json")).cover_letter == "Frozen words."


# ---------------------------------------------------------------- HTTP

def _job_id(client, sample_profile):
    from tests.test_resume_export_routes import accepted_job
    return accepted_job(client, sample_profile)


def test_routes_write_print_download_and_remove_a_letter(client, sample_profile):
    job_id = _job_id(client, sample_profile)
    page = client.get(f"/jobs/{job_id}?step=questions").text
    assert 'id="cover-letter"' in page and "Write the letter yourself" in page
    assert client.get(f"/jobs/{job_id}/cover-letter/print").status_code == 400  # nothing to print yet

    empty = client.post(f"/jobs/{job_id}/cover-letter", data={"text": " "}, follow_redirects=False)
    assert empty.status_code == 422 and "The letter is empty" in empty.text
    saved = client.post(f"/jobs/{job_id}/cover-letter", data={"text": "First paragraph.\n\nSecond paragraph."},
                        follow_redirects=False)
    assert saved.status_code == 303 and saved.headers["location"].endswith("msg=cover_letter_saved#cover-letter")
    assert "Cover letter saved." in client.get(saved.headers["location"]).text

    printed = client.get(f"/jobs/{job_id}/cover-letter/print")
    assert printed.status_code == 200 and "cover-letter-document" in printed.text
    assert 'data-print-resume' in printed.text and 'Print / save PDF</button>' in printed.text
    assert "Dear " in printed.text and "<p>Second paragraph.</p>" in printed.text and "Sincerely," in printed.text
    assert sample_profile["contact"]["email"] in printed.text and "· Cover letter</title>" in printed.text

    response = client.get(f"/jobs/{job_id}/cover-letter/docx")
    assert response.status_code == 200 and "cover%20letter.docx" in response.headers["content-disposition"]
    paragraphs = [p.text for p in Document(BytesIO(response.content)).paragraphs]
    assert paragraphs[0] == sample_profile["contact"]["name"]
    assert ["First paragraph.", "Second paragraph.", "Sincerely,", sample_profile["contact"]["name"]] == paragraphs[-4:]

    review = client.get(f"/jobs/{job_id}?step=review").text
    assert "Cover letter included" in review
    removed = client.post(f"/jobs/{job_id}/cover-letter/remove", follow_redirects=False)
    assert removed.status_code == 303
    assert "No cover letter (optional)" in client.get(f"/jobs/{job_id}?step=review").text


def test_ai_draft_accept_and_discard_over_http(ai_app_client, sample_profile, fake_ai):
    client = ai_app_client
    job_id = _job_id(client, sample_profile)
    approval(client)  # free models share only what the user approved
    with client.app.state.session_factory() as session:
        bullet = profile_service.get_candidate(session).profile.experience[0].bullets[0]
    fake_ai.push(reply(f"I am applying.\n\n{bullet.text}", [bullet.id]))
    drafted_page = client.post(f"/jobs/{job_id}/cover-letter/draft", follow_redirects=True)
    assert "AI draft" in drafted_page.text and bullet.text in drafted_page.text
    draft_token = re.search(r'name="draft_token" value="([^"]+)"', drafted_page.text).group(1)
    stale = client.post(f"/jobs/{job_id}/cover-letter/draft/accept", data={"draft_token": "x"}, follow_redirects=False)
    assert stale.status_code == 409
    accepted = client.post(f"/jobs/{job_id}/cover-letter/draft/accept", data={"draft_token": draft_token},
                           follow_redirects=False)
    assert accepted.status_code == 303
    with client.app.state.session_factory() as session:
        application = session.get(Job, job_id).application
        assert application.cover_letter.origin == "ai_draft" and application.cover_letter_draft is None
    assert client.post(f"/jobs/{job_id}/cover-letter/draft/discard", follow_redirects=False).status_code == 409


def test_submitted_letter_prints_and_downloads_as_frozen(client, sample_profile):
    from tests.test_resume_export_routes import choose

    job_id = _job_id(client, sample_profile)
    assert choose(client, job_id, "modern").status_code == 303
    with client.app.state.session_factory() as session:
        job = session.get(Job, job_id)
        candidate = profile_service.get_candidate(session)
        approve(session, job, candidate)
        record_applied(session, job, candidate)  # no letter in this package
    assert client.get(f"/jobs/{job_id}/snapshots/1/cover-letter").status_code == 404
    assert "Cover letter as submitted" not in client.get(f"/jobs/{job_id}/snapshots/1").text

    with client.app.state.session_factory() as session:
        job = session.get(Job, job_id)
        candidate = profile_service.get_candidate(session)
        job.application.tracking_status = "saved"  # record a second submission of the same job for the test
        session.commit()
        svc.save_letter(session, job, "Frozen paragraph.")
        approve(session, job, candidate)
        record_applied(session, job, candidate)
        svc.save_letter(session, job, "Later paragraph.")
        edited = copy.deepcopy(sample_profile)
        edited["contact"]["name"] = "Changed Current Name"
        profile_service.save(session, edited)

    snapshot_page = client.get(f"/jobs/{job_id}/snapshots/2").text
    assert "Cover letter as submitted" in snapshot_page and "Frozen paragraph." in snapshot_page
    printed = client.get(f"/jobs/{job_id}/snapshots/2/cover-letter")
    assert printed.status_code == 200 and "<p>Frozen paragraph.</p>" in printed.text
    assert "resume-layout-modern" in printed.text and sample_profile["contact"]["name"] in printed.text
    assert "Changed Current Name" not in printed.text and "Later paragraph." not in printed.text
    response = client.get(f"/jobs/{job_id}/snapshots/2/cover-letter.docx")
    document = Document(BytesIO(response.content))
    assert "Frozen paragraph." in [p.text for p in document.paragraphs]
    assert document.styles["Normal"].font.name == "Calibri"  # the snapshot's Modern layout


def test_a_cited_entry_vouches_for_its_own_dates_only(candidate, job):
    dated = f"From June 2024 to August 2024 I worked as an intern.\n\n{BULLET}."
    assert svc.validate_letter(candidate.profile, job, dated, ["exp-1-b1"]) == []  # the bullet's entry dates
    assert svc.validate_letter(candidate.profile, job, dated, ["exp-1"]) != []  # the entry's own text has no 1,200
    assert svc.validate_letter(candidate.profile, job, f"In 2019 I started coding.\n\n{BULLET}.", ["exp-1-b1"])
    # Another entry's dates are not this source's facts.
    assert svc.validate_letter(candidate.profile, job, f"Since November 2023 I built bots.\n\n{BULLET}.",
                               ["exp-1-b1"])


def test_entry_dates_vouch_for_their_years_not_their_month_digits(candidate, job):
    # exp-1 runs 2024-06 to 2024-08: June and August, not the numbers 6 and 8.
    for claim in ("I worked with a team of 6.", "I cut build times by 8%."):
        assert svc.validate_letter(candidate.profile, job, f"{claim}\n\n{BULLET}.", ["exp-1-b1"])


def test_autofill_makes_the_word_letter_from_the_fill_data_it_planned_with():
    from app.services import autofill
    from app.services.autofill_plan import FillData

    letter = svc.document({"name": "Jordan Example", "email": "jordan@example.com"}, "Example Robotics",
                          "I would like to build robots with you.\n\nThank you.")
    data = FillData(contact={}, answers=(), resume_file=None, cover_letter_file="Letter.docx", cover_letter=letter)
    built = Document(BytesIO(autofill.cover_letter_docx(data, layout="classic")))
    assert "I would like to build robots with you." in [p.text for p in built.paragraphs]
    with pytest.raises(autofill.AutofillError):
        autofill.cover_letter_docx(FillData(contact={}, answers=(), resume_file=None), layout="classic")
