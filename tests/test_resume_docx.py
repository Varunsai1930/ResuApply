"""Word export keeps selected facts editable, without leaking omitted content."""

from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.schemas.resume import ResumeContent
from app.services import resume
from app.services import resume_docx
from app.services.resume_docx import ResumeDocxError, export_resume


NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "rels": "http://schemas.openxmlformats.org/package/2006/relationships"}


@pytest.fixture
def rendered():
    profile = {
        "contact": {"name": "Zoë & 李 <Example>", "email": "zoe@example.test", "phone": "+1 555 0100",
                    "location": "São Paulo", "links": {"linkedin": "https://example.test/in/zoe?a=1&b=2",
                    "github": "https://example.test/zoe", "portfolio": "javascript:alert('no')"}},
        "summary": "ORIGINAL SUMMARY THAT WAS NOT ACCEPTED",
        "experience": [
            {"id": "exp-1", "title": "Engineer", "organization": "A & B <Labs>", "location": "Remote",
             "start": "2020-03", "end": "present", "technologies": ["Python", "SQL"],
             "bullets": [{"id": "exp-1-b1", "text": "Original profile bullet omitted from this export"},
                         {"id": "exp-1-b2", "text": "Unselected bullet secret"}]},
            {"id": "exp-2", "title": "Unselected employer secret", "organization": "Hidden Ltd", "bullets": []}],
        "projects": [{"id": "proj-1", "name": "Telemetry", "role": "Maintainer", "link": "https://example.test/telemetry",
                      "start": "2022", "end": "2023", "technologies": ["SQLite"],
                      "bullets": [{"id": "proj-1-b1", "text": "Original project wording"}]}],
        "education": [{"id": "edu-1", "institution": "Example University", "degree": "BSc", "field": "Computer Science",
                       "start": "2016", "end": "2020", "gpa": "3.8/4.0",
                       "bullets": [{"id": "edu-1-b1", "text": "Completed an honors thesis."}]},
                      {"id": "edu-2", "institution": "Unselected school secret", "bullets": []}],
        "skills": [{"name": "Python"}, {"name": "SQL"}, {"name": "Unselected skill secret"}],
        "certifications": [{"id": "cert-1", "name": "Professional Certificate", "issuer": "Example Board", "date": "2024-02"},
                           {"id": "cert-2", "name": "Unselected certificate secret"}],
    }
    content = ResumeContent.model_validate({
        "summary": {"text": "Accepted summary with naïve résumé and 🚀.", "sources": ["summary"]},
        "experience": [{"entry": "exp-1", "bullets": [{"text": "Accepted wording: improved café tools.\nPreserved second line.", "sources": ["exp-1-b1"]}]}],
        "projects": [{"entry": "proj-1", "bullets": [{"text": "Maintained the selected telemetry project.", "sources": ["proj-1-b1"]}]}],
        "education": ["edu-1"], "skills": ["SQL", "Python"], "certifications": ["cert-1"],
    })
    return resume.render_context(profile, content)


def xml_parts(data):
    with ZipFile(BytesIO(data)) as package:
        return {name: ET.fromstring(package.read(name)) for name in
                ("word/document.xml", "word/numbering.xml", "word/styles.xml", "word/_rels/document.xml.rels", "docProps/core.xml")}


@pytest.mark.parametrize("layout", ["classic", "modern", "compact"])
def test_export_contains_only_the_selected_canonical_facts(rendered, layout):
    before = deepcopy(rendered)
    data = export_resume(rendered, layout)
    document = Document(BytesIO(data))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for selected in ("Zoë & 李 <Example>", "zoe@example.test", "+1 555 0100", "São Paulo",
                     "Accepted summary with naïve résumé and 🚀.", "Engineer — A & B <Labs>",
                     "2020-03 – Present", "Technologies: Python, SQL",
                     "Accepted wording: improved café tools.\nPreserved second line.",
                     "Telemetry — Maintainer", "2022 – 2023", "https://example.test/telemetry",
                     "BSc, Computer Science", "2016 – 2020", "GPA: 3.8/4.0",
                     "Completed an honors thesis.", "SQL, Python",
                     "Professional Certificate — Example Board (2024-02)"):
        assert selected in text
    for omitted in ("ORIGINAL SUMMARY", "Original profile bullet", "Original project wording", "Unselected",
                    "exp-1-b1", "proj-1-b1"):
        assert omitted not in text
    assert rendered == before
    core = ET.tostring(xml_parts(data)["docProps/core.xml"], encoding="unicode")
    assert "Zoë" not in core and "A & B" not in core and "python-docx" not in core
    assert not document.core_properties.author and not document.core_properties.last_modified_by
    assert document.core_properties.comments == ""


def test_export_uses_real_numbered_bullets_and_safe_external_hyperlinks(rendered):
    parts = xml_parts(export_resume(rendered))
    body, numbering, relationships = (parts[name] for name in
                                     ("word/document.xml", "word/numbering.xml", "word/_rels/document.xml.rels"))
    bullets = [
        paragraph for paragraph in body.findall(".//w:p", NS) if paragraph.find("w:pPr/w:numPr", NS) is not None]
    assert len(bullets) == 4  # experience, project, education, certification
    for bullet in bullets:
        assert not "".join(bullet.itertext()).startswith(("• ", "- "))
        num_id = bullet.find("w:pPr/w:numPr/w:numId", NS).get(f"{{{NS['w']}}}val")
        num = next(value for value in numbering.findall("w:num", NS) if value.get(f"{{{NS['w']}}}numId") == num_id)
        abstract_id = num.find("w:abstractNumId", NS).get(f"{{{NS['w']}}}val")
        abstract = next(value for value in numbering.findall("w:abstractNum", NS)
                        if value.get(f"{{{NS['w']}}}abstractNumId") == abstract_id)
        assert abstract.find("w:lvl/w:numFmt", NS).get(f"{{{NS['w']}}}val") == "bullet"
    external = [item for item in relationships if item.get("TargetMode") == "External"]
    assert {item.get("Target") for item in external} == {
        "https://example.test/in/zoe?a=1&b=2", "https://example.test/zoe", "https://example.test/telemetry"}
    assert len(body.findall(".//w:hyperlink", NS)) == 3
    assert "javascript:alert('no')" in "".join(body.itertext())


@pytest.mark.parametrize("url", ["file:///tmp/private", "javascript:alert(1)", "https://", "https://[broken",
                                 "https://example.test:99999/", "https://example.test:invalid/", "https://exam ple.test/"])
def test_non_web_or_malformed_links_remain_plain_text(rendered, url):
    rendered["contact"]["links"] = {"portfolio": url}
    rendered["projects"] = []
    parts = xml_parts(export_resume(rendered))
    assert parts["word/document.xml"].find(".//w:hyperlink", NS) is None
    assert url in "".join(parts["word/document.xml"].itertext())


@pytest.mark.parametrize("layout,font,size,heading,leading", [
    ("classic", "Arial", 10.5, 12, 1.40), ("modern", "Calibri", 10.5, 11, 1.45),
    ("compact", "Arial", 10, 10.5, 1.32)])
def test_layouts_are_a4_editable_and_can_flow_across_pages(rendered, layout, font, size, heading, leading):
    document = Document(BytesIO(export_resume(rendered, layout)))
    assert round(document.sections[0].page_width.mm) == 210
    assert round(document.sections[0].page_height.mm) == 297
    assert document.styles["Normal"].font.name == font
    assert document.styles["Normal"].font.size.pt == size
    assert document.styles["Resume Section"].font.size.pt == heading
    assert document.paragraphs[0].style.name == "Title"
    assert str(document.styles["Title"].font.color.rgb) == "000000"
    assert document.styles["Title"].element.find("w:pPr/w:pBdr", NS) is None
    assert document.paragraphs[0]._p.find("w:pPr/w:pBdr", NS) is None
    # Word stores multiple line spacing in integer 1/240 units.
    assert document.styles["Normal"].paragraph_format.line_spacing == pytest.approx(leading, abs=1 / 240)
    assert document.styles["Resume Section"].paragraph_format.keep_with_next
    entries = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "Resume Entry"]
    bullets = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "Resume Bullet"]
    assert all(paragraph.paragraph_format.keep_with_next for paragraph in entries)
    assert all(not paragraph.style.paragraph_format.keep_with_next for paragraph in bullets)
    assert all(not paragraph.style.paragraph_format.keep_together for paragraph in bullets)
    assert not document.tables


def test_empty_optional_sections_and_single_dates_do_not_create_placeholders():
    document = Document(BytesIO(export_resume({"contact": {"name": "Example"}, "experience": [
        {"organization": "Lab", "start": "2024", "end": "", "bullets": []}]})))
    assert [paragraph.text for paragraph in document.paragraphs] == ["Example", "Experience", "Lab", "2024"]
    assert document.paragraphs[-1].paragraph_format.keep_with_next is False


def test_minimal_modern_name_has_no_decorative_rule():
    document = xml_parts(export_resume({"contact": {"name": "Example"}}, "modern"))["word/document.xml"]
    assert document.find(".//w:pBdr", NS) is None


@pytest.mark.parametrize("layout", ["classic", "modern", "compact"])
def test_template_title_borders_are_removed_from_the_export(monkeypatch, layout):
    template = Document()
    borders = OxmlElement("w:pBdr")
    underline = OxmlElement("w:bottom")
    underline.set(qn("w:val"), "single")
    underline.set(qn("w:color"), "4F81BD")
    borders.append(underline)
    template.styles["Title"].element.get_or_add_pPr().append(borders)
    monkeypatch.setattr(resume_docx, "Document", lambda: template)
    parts = xml_parts(export_resume({"contact": {"name": "Example"}}, layout))
    title_style = next(style for style in parts["word/styles.xml"]
                       if style.get(qn("w:styleId")) == "Title")
    assert title_style.find("w:pPr/w:pBdr", NS) is None
    assert parts["word/document.xml"].find(".//w:pBdr", NS) is None


def test_long_text_is_preserved_without_fixed_height_containers(rendered):
    long_text = "A very long accepted claim with international text résumé. " * 1000
    rendered["experience"][0]["bullets"][0]["text"] = long_text
    data = export_resume(rendered, "compact")
    document = Document(BytesIO(data))
    assert long_text in [paragraph.text for paragraph in document.paragraphs]
    body = xml_parts(data)["word/document.xml"]
    assert body.find(".//w:trHeight", NS) is None
    assert body.find(".//w:txbxContent", NS) is None


@pytest.mark.parametrize("bad", ["\x00", "\x01", "\x0b", "\ud800", "\ufffe", "\uffff"])
def test_unrepresentable_accepted_text_produces_a_readable_error(rendered, bad):
    rendered["summary"]["text"] += bad
    with pytest.raises(ResumeDocxError, match="Word cannot read"):
        export_resume(rendered)


@pytest.mark.parametrize("damaged", [None, {"contact": []}, {"contact": {"name": 1}},
    {"contact": {"name": "Example"}, "experience": "not a list"},
    {"contact": {"name": "Example"}, "skills": [None]},
    {"contact": {"name": "Example"}, "summary": "not a claim"},
    {"contact": {"name": "Example"}, "education": [{"institution": "Example", "bullets": ["not a bullet"]}]}])
def test_damaged_context_produces_a_readable_error(damaged):
    with pytest.raises(ResumeDocxError, match="Resume"):
        export_resume(damaged)


@pytest.mark.parametrize("layout", ["unknown", "../../evil", "", None, ["classic"]])
def test_invalid_layout_produces_a_readable_error(rendered, layout):
    with pytest.raises(ResumeDocxError, match="valid resume layout"):
        export_resume(rendered, layout)
