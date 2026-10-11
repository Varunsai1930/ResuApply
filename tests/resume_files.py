"""Small resume files built in memory for the import tests (synthetic text only)."""

from __future__ import annotations

import io


def pdf_bytes(lines: list[str]) -> bytes:
    """A one-page text PDF with each line on its own line, in Helvetica."""
    def escape(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    body = "BT /F1 11 Tf 14 TL 72 760 Td " + " ".join(f"({escape(line)}) Tj T*" for line in lines) + " ET"
    stream = body.encode("cp1252")  # WinAnsi, as the font declares: bullets and dashes included
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, content in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + content + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def encrypted_pdf_bytes(lines: list[str], password: str = "secret") -> bytes:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(pdf_bytes(lines))))
    writer.encrypt(password)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def blank_pdf_bytes() -> bytes:
    """A PDF with a page but no text, as a scanned resume is to a text reader."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def docx_bytes(lines: list[str], table: list[list[str]] | None = None) -> bytes:
    from docx import Document

    document = Document()
    for line in lines:
        document.add_paragraph(line)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, text in enumerate(row):
                grid.cell(r, c).text = text
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- laid-out PDFs

# Characters outside WinAnsi that real resumes' PDFs extract to, each drawn with a spare code the
# font's ToUnicode map gives back: ligatures (LaTeX), bullet glyphs (Word's Symbol font, design tools).
EXTRA_CHARS = "ﬁﬂﬀﬃ▪●◦➢✓"


def _cmap() -> bytes:
    pairs = [(code, bytes([code]).decode("cp1252")) for code in range(0x20, 0x100)
             if code not in (0x81, 0x8D, 0x8F, 0x90, 0x9D)]
    pairs += [(i + 1, ch) for i, ch in enumerate(EXTRA_CHARS)]
    body = "\n".join(f"<{code:02X}> <{ord(ch):04X}>" for code, ch in pairs)
    return ("/CIDInit /ProcSet findresource begin 12 dict begin begincmap /CMapName /ResuTest def "
            "1 begincodespacerange <00> <FF> endcodespacerange\n"
            f"{len(pairs)} beginbfchar\n{body}\nendbfchar endcmap CMapName currentdict /CMap defineresource pop end end"
            ).encode()


def _encode(text: str) -> bytes:
    out = bytearray()
    for ch in text:
        if ch in EXTRA_CHARS:
            out.append(EXTRA_CHARS.index(ch) + 1)
        else:
            out += ch.encode("cp1252")
    return bytes(out).replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


def layout_pdf_bytes(pages: list[list[tuple[float, float, str]]], size: float = 10) -> bytes:
    """A PDF whose pages draw each ``(x, y, text)`` in the order given, as a layout tool would.

    The drawing order is what a text reader follows, so a two-column page can be drawn column
    by column or line by line across both columns, as different resume builders do.
    """
    objects: list[bytes] = [b"", b""]  # catalog and page tree, filled in below
    font = len(objects) + 1
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding "
                   + f"/ToUnicode {font + 1} 0 R >>".encode())
    cmap = _cmap()
    objects.append(b"<< /Length " + str(len(cmap)).encode() + b" >>\nstream\n" + cmap + b"\nendstream")
    kids = []
    for page in pages:
        ops = b"".join(b"BT /F1 " + str(size).encode() + b" Tf " + f"{x:.1f} {y:.1f} Td (".encode() + _encode(text)
                       + b") Tj ET\n" for x, y, text in page)
        objects.append(b"<< /Length " + str(len(ops)).encode() + b" >>\nstream\n" + ops + b"endstream")
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {len(objects)} 0 R "
                       f"/Resources << /Font << /F1 {font} 0 R >> >> >>".encode())
        kids.append(len(objects))
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>".encode()
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, content in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + content + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def column(lines: list[str], x: float = 50, top: float = 750, leading: float = 14) -> list[tuple[float, float, str]]:
    """Lines one under another from ``top``, as one column draws them."""
    return [(x, top - i * leading, line) for i, line in enumerate(lines)]
