"""Corrupt compressed Word parts fail safely without changing a saved profile."""

import io
import zipfile

import pytest

from app.services.resume_import import ResumeImportError, extract_text
from tests.resume_files import docx_bytes
from tests.test_damaged_docx_routes import _profiles
from tests.test_routes import profile_form, review_and_save


LINES = ["Jordan Example", "jordan@example.com", "Software developer building reliable Python applications."]
COMPRESSION = pytest.mark.parametrize("compression", [zipfile.ZIP_BZIP2, zipfile.ZIP_LZMA], ids=["bzip2", "lzma"])


def _compressed_docx(compression, *, damaged=False):
    """Keep the ZIP directory intact while corrupting only a document's compressed data."""
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(docx_bytes(LINES))) as source, zipfile.ZipFile(output, "w") as archive:
        for name in source.namelist():
            method = compression if name == "word/document.xml" else zipfile.ZIP_DEFLATED
            archive.writestr(name, source.read(name), compress_type=method)
    data = bytearray(output.getvalue())
    if damaged:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            part = archive.getinfo("word/document.xml")
        start = part.header_offset + 30 + len(part.filename.encode()) + len(part.extra)
        data[start:start + 10] = b"\x00" * 10
    return bytes(data)


def _mismatched_header_type():
    from docx import Document

    document = Document()
    document.sections[0].header.paragraphs[0].text = LINES[0]
    for line in LINES[1:]:
        document.add_paragraph(line)
    source = io.BytesIO()
    document.save(source)
    output = io.BytesIO()
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as changed:
        for name in archive.namelist():
            data = archive.read(name)
            if name == "[Content_Types].xml":
                data = data.replace(b"application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml",
                                    b"application/xml")
            changed.writestr(name, data)
    return output.getvalue()


@COMPRESSION
def test_valid_alternative_compression_remains_readable(compression):
    assert extract_text(_compressed_docx(compression)).splitlines() == LINES


@COMPRESSION
def test_corrupt_alternative_compression_is_a_recoverable_import_error(compression):
    with pytest.raises(ResumeImportError, match="This Word file couldn't be read"):
        extract_text(_compressed_docx(compression, damaged=True))


def test_a_mismatched_word_header_type_is_a_recoverable_import_error():
    with pytest.raises(ResumeImportError, match="This Word file couldn't be read"):
        extract_text(_mismatched_header_type())


@pytest.mark.parametrize("make", [
    pytest.param(lambda: _compressed_docx(zipfile.ZIP_BZIP2, damaged=True), id="bzip2"),
    pytest.param(lambda: _compressed_docx(zipfile.ZIP_LZMA, damaged=True), id="lzma"),
    pytest.param(_mismatched_header_type, id="header-content-type"),
])
@pytest.mark.parametrize("saved_profile", [False, True], ids=["new-profile", "saved-profile"])
def test_damaged_word_upload_keeps_the_profile_and_form_available(client, sample_profile, make, saved_profile):
    if saved_profile:
        _, saved = review_and_save(client, profile_form(sample_profile))
        assert saved.status_code == 303
    before = _profiles(client)

    response = client.post("/profile/import", files={"file": (
        "resume.docx", make(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )})

    assert response.status_code == 422
    assert "This Word file couldn&#39;t be read" in response.text
    assert 'name="file"' in response.text
    assert _profiles(client) == before
    assert client.get("/profile").status_code == 200
